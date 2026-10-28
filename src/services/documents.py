"""Turn the Markdown that the agents write into .docx and PDF files."""
import io
import re

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

_BOLD = re.compile(r"(\*\*[^*]+\*\*)")
_REPLACE = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-",
            "\u2022": "-", "\u2026": "...", "\u00a0": " ", "\u2192": "->"}


def _lines(md: str):
    """Yield (kind, text): h1/h2/h3, bullet, para."""
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            yield f"h{len(m.group(1))}", m.group(2).strip()
        elif re.match(r"^\s*[-*+]\s+", line):
            yield "bullet", re.sub(r"^\s*[-*+]\s+", "", line)
        else:
            yield "para", line.strip()


def to_docx(md: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for kind, text in _lines(md):
        if kind.startswith("h"):
            doc.add_heading(text.replace("**", ""), level=int(kind[1]))
            continue
        para = doc.add_paragraph(style="List Bullet" if kind == "bullet" else None)
        for chunk in _BOLD.split(text):
            if chunk.startswith("**") and chunk.endswith("**"):
                para.add_run(chunk[2:-2]).bold = True
            elif chunk:
                para.add_run(chunk)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _latin(text: str) -> str:
    for a, b in _REPLACE.items():
        text = text.replace(a, b)
    text = text.encode("latin-1", "replace").decode("latin-1")
    # break very long unbroken strings (URLs) so the PDF layout engine can wrap them
    return re.sub(r"(\S{60})(?=\S)", r"\1 ", text)


def to_pdf(md: str) -> bytes:
    pdf = FPDF(format="Letter")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, 16)
    pdf.add_page()
    sizes = {"h1": 18, "h2": 12.5, "h3": 11}
    for kind, text in _lines(md):
        text = _latin(text)
        pdf.set_x(pdf.l_margin)
        if kind in sizes:
            pdf.ln(2 if kind != "h1" else 0)
            pdf.set_font("Helvetica", "B", sizes[kind])
            pdf.multi_cell(0, sizes[kind] * 0.5, text.replace("**", ""), new_x="LMARGIN", new_y="NEXT")
            if kind == "h2":
                y = pdf.get_y()
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(1.5)
            continue
        pdf.set_font("Helvetica", "", 10)
        if kind == "bullet":
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(0, 5, "- " + text, markdown=True, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.multi_cell(0, 5, text, markdown=True, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
